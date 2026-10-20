from flask import Flask, render_template, request, jsonify, send_file
from flask_cors import CORS
import pandas as pd
import os
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from fpdf import FPDF

# ¡Esta es la línea que Gunicorn no encontraba!
app = Flask(__name__)
# Habilitar CORS para que Netlify pueda hablar con Render
CORS(app)

@app.route('/')
def inicio():
    return render_template('liquidador.html')

@app.route('/guardar_historial', methods=['POST'])
def guardar_historial():
    datos = request.json
    datos['Fecha_Cotizacion'] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    datos_ordenados = {"Fecha": datos['Fecha_Cotizacion'], **datos}
    
    df_nuevo = pd.DataFrame([datos_ordenados])
    nombre_archivo = "Historial_Importaciones.xlsx"
    
    if os.path.exists(nombre_archivo):
        df_existente = pd.read_excel(nombre_archivo)
        df_final = pd.concat([df_existente, df_nuevo], ignore_index=True)
    else:
        df_final = df_nuevo
        
    df_final.to_excel(nombre_archivo, index=False)
    
    wb = load_workbook(nombre_archivo)
    ws = wb.active
    fondo_azul = PatternFill(start_color="0D1B2A", end_color="0D1B2A", fill_type="solid")
    letras_blancas = Font(color="FFFFFF", bold=True)
    
    for celda in ws[1]: 
        celda.fill = fondo_azul
        celda.font = letras_blancas
        celda.alignment = Alignment(horizontal="center")
        
    for col in ws.columns:
        max_length = 0
        letra_columna = col[0].column_letter
        for celda in col:
            try:
                if len(str(celda.value)) > max_length:
                    max_length = len(str(celda.value))
            except: pass
        ws.column_dimensions[letra_columna].width = max_length + 2

    wb.save(nombre_archivo)
    return jsonify({"estatus": "ok", "mensaje": "Historial guardado con Fecha y Formato PRO"})

@app.route('/descargar_pdf', methods=['POST'])
def descargar_pdf():
    datos = request.json
    fecha_actual = datetime.now().strftime('%Y-%m-%d %H:%M')
    
    pdf = FPDF()
    pdf.add_page()
    
    pdf.set_fill_color(13, 27, 42)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 15, "LIQUIDADOR PRO COLOMBIA - REPORTE OFICIAL", ln=True, align="C", fill=True)
    
    pdf.ln(10)
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 10, f"Fecha del Calculo: {fecha_actual}", ln=True)
    pdf.ln(5)
    
    pdf.set_font("Arial", "", 12)
    pdf.cell(90, 10, f"Valor FOB: $ {datos['FOB_USD']} USD", ln=False)
    pdf.cell(90, 10, f"TRM Aplicada: {datos['TRM']} COP", ln=True)
    
    pdf.cell(90, 10, f"Flete: $ {datos['Flete_USD']} USD", ln=False)
    pdf.cell(90, 10, f"Seguro: $ {datos['Seguro_USD']} USD", ln=True)
    
    pdf.ln(15)
    
    pdf.set_fill_color(255, 183, 3) 
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 20, f"COSTO LANDED TOTAL: {datos['Costo_Landed_COP']}", ln=True, align="C", fill=True)
    
    ruta_pdf = "Cotizacion_Landed.pdf"
    pdf.output(ruta_pdf)
    
    return send_file(ruta_pdf, as_attachment=True, download_name="Liquidacion_Oficial.pdf", mimetype='application/pdf')

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port)
